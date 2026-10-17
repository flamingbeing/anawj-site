# Check a recalculated copy of a generated sheet: every right-hand total and every count under the names must equal what is
# worked out independently from the day cells. Usage (needs LibreOffice and openpyxl):
#   node nuh-junior-monthly/tools/sheetgen-demo.mjs 2026-12 14          # writes /tmp/sheetgen-2026-12-14.xlsx (made-up people)
#   soffice --headless --convert-to xlsx:"Calc MS Excel 2007 XML" --outdir /tmp/recalc /tmp/sheetgen-2026-12-14.xlsx
#   python3 -I nuh-junior-monthly/tools/sheetgen-verify.py /tmp/recalc/sheetgen-2026-12-14.xlsx
import sys,openpyxl
from openpyxl.utils import get_column_letter as L
path=sys.argv[1]
wb=openpyxl.load_workbook(path,data_only=True); ws=wb.worksheets[0]
# find layout from the file itself
hdr={}
for c in range(1,ws.max_column+1):
    v=ws.cell(5,c).value
    if v in ('WE','Fri','Sat','Sun','WD','Total','AL') and c>5 and c not in hdr.values():
        hdr.setdefault(v,c)
n=sum(1 for c in range(4,40) if isinstance(ws.cell(6,c).value,(int,float)))
first=4; last=first+n-1; nxt=last+1
# right block columns by label order
cols={}
for c in range(nxt+1,ws.max_column+1):
    v=ws.cell(5,c).value
    if v is not None: cols[v]=c
people=[]
r=7
while ws.cell(r,3).value not in (None,''):
    people.append(r); r+=1
errs=0; checked=0
def cnt(r,code,upto): return sum(1 for c in range(first,upto+1) if ws.cell(r,c).value==code)
for r in people:
    wd=lambda c: ws.cell(5,c).value
    def pts(day,upto):
        return sum(1 for c in range(first,upto+1) if wd(c)==day and ws.cell(r,c).value in ('R1','R2','R3','S1','S1B'))
    exp={
      'WE': sum(cnt(r,k,last) for k in ('WE','NP','PH','RD','X','XX','X/NP')),
      'Fri':1.5*pts('Fri',last),'Sat':2*pts('Sat',last),
      'Sun':1.5*pts('Sun',last),   # no holidays in this fake
      'WD':sum(pts(d,nxt) for d in ('Mon','Tue','Wed','Thu')),
      'R1B':cnt(r,'R1B',last),'C3':cnt(r,'C3',last),'OGN':cnt(r,'OGN',last)*0.5,'OGC':cnt(r,'OGC',last),'S1':cnt(r,'S1',last),'S1B':cnt(r,'S1B',last),
      'AL':sum(cnt(r,k,nxt) for k in ('AL','SL','ML','PL','FCL','CCL','MC','HL','NS')),
    }
    exp['Total']=exp['Fri']+exp['Sat']+exp['Sun']+exp['WD']+exp['R1B']+exp['C3']+exp['OGN']+exp['OGC']
    for k,v in exp.items():
        got=ws.cell(r,cols[k]).value; checked+=1
        if got is None or abs(float(got)-v)>1e-9: errs+=1; print('MISMATCH row',r,k,'got',got,'expected',v)
# bottom counts
rr=people[-1]+2; codes=[]
while ws.cell(rr,3).value not in (None,''):
    code=ws.cell(rr,3).value
    for c in range(first,last+1):
        exp=sum(1 for r in people if ws.cell(r,c).value==code); got=ws.cell(rr,c).value; checked+=1
        if got!=exp: errs+=1; print('MISMATCH count',code,L(c),got,exp)
    rr+=1
# error cells anywhere
bad=[c.coordinate for row in ws.iter_rows() for c in row if isinstance(c.value,str) and c.value.startswith('#')]
print(f'{path}: {len(people)} people, {n} days; checked {checked} values; mismatches {errs}; error cells {bad[:5]}')
sys.exit(1 if errs or bad else 0)
